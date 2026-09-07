# -*- coding: utf-8 -*-
"""
Унифицированная SQLite база данных для всех процедур.
Одна БД с таблицей permits для разделения по типам.
"""
import csv
import json
import logging
import os
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import tkinter.filedialog as filedialog
    import tkinter.messagebox as messagebox
except ImportError:
    messagebox = None
    filedialog = None

logger = logging.getLogger(__name__)


class UnifiedPermitDB:
    """Единая база данных для всех процедур пропусков."""

    def __init__(self, db_path: Optional[str] = None):
        if db_path is None:
            db_path = get_db_path()
        self._path = str(db_path)
        Path(self._path).parent.mkdir(parents=True, exist_ok=True)
        self._conn = sqlite3.connect(self._path, check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        self._init_schema()

    def _init_schema(self) -> None:
        """Инициализация таблиц и индексов."""
        with self._conn:
            self._conn.execute("""
                CREATE TABLE IF NOT EXISTS permits (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    procedure_code TEXT NOT NULL,    -- '14.3', '14.5', '19.17.1'
                    search_key TEXT NOT NULL,        -- ключ поиска (ФИО или организация)
                    search_type TEXT NOT NULL,       -- 'fio' или 'org'
                    data TEXT NOT NULL,              -- JSON с полными данными
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
            self._conn.execute("""
                CREATE INDEX IF NOT EXISTS idx_permits_search
                ON permits(search_key, search_type, procedure_code)
            """)
            self._conn.execute("""
                CREATE INDEX IF NOT EXISTS idx_permits_proc
                ON permits(procedure_code, created_at DESC)
            """)

    def close(self) -> None:
        """Закрытие соединения с БД."""
        try:
            self._conn.close()
        except Exception as e:
            logger.warning(f"Ошибка при закрытии БД: {e}")

    # ---- Чтение ----
    def load(self, procedure_code: Optional[str] = None, search_type: Optional[str] = None) -> List[Dict[str, Any]]:
        """Все записи (опционально фильтруя по процедуре и типу поиска)."""
        sql = "SELECT data, search_key, procedure_code FROM permits WHERE 1=1"
        params: List[Any] = []
        if procedure_code:
            sql += " AND procedure_code = ?"
            params.append(procedure_code)
        if search_type:
            sql += " AND search_type = ?"
            params.append(search_type)
        sql += " ORDER BY id DESC"

        rows = self._conn.execute(sql, params).fetchall()
        out = []
        for row in rows:
            try:
                rec = json.loads(row["data"])
            except json.JSONDecodeError as e:
                logger.error(f"Ошибка декодирования JSON для ключа {row['search_key']}: {e}")
                rec = {}
            rec["_search_key"] = row["search_key"]
            rec["_procedure_code"] = row["procedure_code"]
            out.append(rec)
        return out

    def find(self, key: str, procedure_code: str, search_type: str) -> Optional[Dict[str, Any]]:
        """Точный поиск по ключу (регистронезависимый)."""
        k = (key or "").strip().lower()
        if not k:
            return None
        row = self._conn.execute(
            "SELECT data FROM permits WHERE search_key=? AND search_type=? AND procedure_code=? ORDER BY id DESC LIMIT 1",
            (k, search_type, procedure_code)
        ).fetchone()
        
        if not row:
            return None
        try:
            return json.loads(row["data"])
        except json.JSONDecodeError as e:
            logger.error(f"Ошибка декодирования JSON при поиске ключа {k}: {e}")
            return None

    def find_like(self, fragment: str, procedure_code: str, search_type: str, limit: int = 50) -> List[str]:
        """Поиск по фрагменту ключа (для автодополнения)."""
        frag = (fragment or "").strip().lower()
        if not frag:
            return []
        rows = self._conn.execute(
            "SELECT DISTINCT search_key FROM permits WHERE search_type=? AND procedure_code=? AND search_key LIKE ? ORDER BY search_key LIMIT ?",
            (search_type, procedure_code, f"%{frag}%", limit)
        ).fetchall()
        return [r["search_key"] for r in rows]

    # ---- Запись ----
    def upsert(self, rec: Dict[str, Any], procedure_code: str, search_type: str) -> None:
        """Добавляет/обновляет запись. search_key извлекается из rec или формируется автоматически."""
        try:
            # Определяем ключ поиска
            if rec.get("search_key"):
                key = str(rec["search_key"]).strip().lower()
            elif search_type == "fio":
                key = " ".join(str(rec.get(k, "")) for k in ("last_name", "first_name", "middle_name") if rec.get(k)).strip().lower()
            else:  # org
                key = str(rec.get("org_short") or rec.get("org_info") or rec.get("org_key") or "").strip().lower()

            if not key:
                logger.warning("Пропущена запись без ключа поиска: %s", rec)
                return

            with self._conn:
                # Удаляем старые записи с таким же ключом
                self._conn.execute(
                    "DELETE FROM permits WHERE search_key=? AND search_type=? AND procedure_code=?",
                    (key, search_type, procedure_code)
                )
                # Вставляем новую
                self._conn.execute(
                    "INSERT INTO permits (procedure_code, search_key, search_type, data, created_at) VALUES (?,?,?,?,?)",
                    (procedure_code, key, search_type, json.dumps(rec, ensure_ascii=False), datetime.now().isoformat())
                )
        except Exception as e:
            logger.exception(f"Ошибка при создании/обновлении записи: {e}")

    def replace_all(self, records: List[Dict[str, Any]], procedure_code: str, search_type: str) -> None:
        """Полная перезапись (миграция)."""
        with self._conn:
            self._conn.execute("DELETE FROM permits WHERE procedure_code=? AND search_type=?", (procedure_code, search_type))
            for rec in records:
                self.upsert(rec, procedure_code, search_type)

    # ---- Экспорт ----
    def export(self, out_path: str, fmt: str = "json", procedure_code: Optional[str] = None) -> Optional[str]:
        """Экспорт в JSON/CSV/Excel."""
        recs = self.load(procedure_code)

        if fmt == "json":
            # Удаляем системные служебные поля при экспорте
            clean_recs = [{k: v for k, v in r.items() if not k.startswith("_")} for r in recs]
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(clean_recs, f, ensure_ascii=False, indent=1)
            return out_path

        # Сбор заголовков без служебных полей
        fieldnames = sorted({k for r in recs for k in r.keys() if not k.startswith("_")}) or ["search_key"]

        if fmt == "csv":
            with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                for r in recs:
                    writer.writerow({k: r.get(k, "") for k in fieldnames})
            return out_path

        if fmt == "excel":
            try:
                from openpyxl import Workbook
                from openpyxl.utils import get_column_letter
            except ImportError:
                logger.error("Модуль openpyxl не установлен.")
                return None

            wb = Workbook()
            ws = wb.active
            ws.title = "Пропуска"
            ws.append(fieldnames)

            for r in recs:
                ws.append([str(r.get(k, "")) for k in fieldnames])

            # Безопасное автотипирование ширины колонок
            for col_idx, field in enumerate(fieldnames, 1):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max(len(str(field)) + 3, 15)

            wb.save(out_path)
            return out_path

        return None

    # ---- Миграция из старых JSON ----
    def migrate_from_json(self, json_path: str, procedure_code: str, search_type: str) -> int:
        """Перенос данных из старого JSON файла."""
        if not os.path.exists(json_path):
            return 0
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                recs = json.load(f)
        except Exception as e:
            logger.error(f"Ошибка чтения JSON при миграции {json_path}: {e}")
            return 0

        if not isinstance(recs, list) or not recs:
            return 0

        migrated = 0
        for rec in recs:
            if isinstance(rec, dict):
                self.upsert(rec, procedure_code, search_type)
                migrated += 1

        # Переименуем старый файл после успешной обработки
        try:
            os.rename(json_path, f"{json_path}.migrated")
        except OSError as e:
            logger.warning(f"Не удалось переименовать файл {json_path}: {e}")

        return migrated


def get_db_path() -> str:
    """Путь к БД рядом с exe/скриптом."""
    if getattr(sys, 'frozen', False):
        base = Path(os.path.dirname(os.path.abspath(sys.argv[0])))
    else:
        base = Path(__file__).parent.parent
    return str(base / "data" / "permits_unified.sqlite")


def export_db_dialog(app, db_path: str, default_name: str) -> None:
    """Универсальный диалог выгрузки БД в JSON/CSV/Excel."""
    if not messagebox or not filedialog:
        logger.error("Tkinter GUI элементы недоступны.")
        return

    if not os.path.exists(db_path):
        messagebox.showinfo("Экспорт БД", "База данных пуста или не создана.", parent=app)
        return

    stamp = datetime.now().strftime("%Y-%m-%d")
    out = filedialog.asksaveasfilename(
        parent=app,
        title="Выгрузить базу пропусков",
        defaultextension=".json",
        initialfile=default_name.replace("XXXX", stamp),
        filetypes=[("JSON", "*.json"), ("Excel", "*.xlsx"), ("CSV", "*.csv")]
    )
    if not out:
        return

    ext = os.path.splitext(out)[1].lower()
    fmt_map = {".xlsx": "excel", ".csv": "csv", ".json": "json"}
    fmt = fmt_map.get(ext, "json")

    try:
        db = UnifiedPermitDB(db_path)
        res = db.export(out, fmt=fmt)
        db.close()

        if fmt == "excel" and res is None:
            messagebox.showerror(
                "Экспорт БД",
                "Не удалось экспортировать в Excel (нужен модуль openpyxl).\n"
                "Установите: pip install openpyxl",
                parent=app,
            )
            return

        messagebox.showinfo("Экспорт БД", f"База выгружена:\n{out}", parent=app)
    except Exception as e:
        logger.exception("Ошибка экспорта")
        messagebox.showerror("Экспорт БД", f"Ошибка: {e}", parent=app)
